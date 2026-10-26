"""강의 슬라이드 생성 (템플릿 디자인 적용).

python3 build_pptx.py <출력 경로>

template/tricentis_base.pptx 의 마스터·레이아웃·테마를 그대로 쓰고, 슬라이드 구성 요소
(제목, 3단 규칙선, 표, 번호 사각형, LAB 화면, 하단 각주와 꼬리말)는 원본 템플릿의
위치·크기·색을 따라 그린다. 프롬프트와 템플릿 문구는 content.js 에서 읽는다.
"""
import json
import math
import os
import subprocess
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else "deck.pptx"
K = json.loads(subprocess.check_output(
    ["node", "-e", "console.log(JSON.stringify(require(process.argv[1])))", os.path.join(HERE, "content.js")]))

# 템플릿에서 쓰는 색
BLUE, TEAL, LIME, GRAY = "004C97", "0E88AE", "D2D755", "4E4F51"
LIGHT, LINE, WHITE, ORANGE = "F1F5F9", "DCE5EF", "FFFFFF", "E8742A"
LATIN, EA, MONO = "Open Sans", "Noto Sans CJK KR", "Consolas"
FOOTER = "Tosca Commander MCP 입문  |  AI 어시스턴트 실습"

prs = Presentation(os.path.join(HERE, "template", "tricentis_base.pptx"))
LAYOUT = {l.name: l for l in prs.slide_layouts}
WARN = []


# ---------- 글자 ----------
def _em(ch):
    o = ord(ch)
    if 0xAC00 <= o <= 0xD7A3 or 0x3000 <= o <= 0x9FFF or 0xFF00 <= o <= 0xFFEF:
        return 1.0
    if ch == " ":
        return 0.28
    if ch in "il.,:;'|!()[]{}":
        return 0.3
    if ch.isupper() or ch in "mwMW@_":
        return 0.68
    return 0.56


def est_lines(text, width_in, size):
    lines = 0
    for part in text.split("\n"):
        em = sum(_em(c) for c in part)
        lines += max(1, math.ceil(em * size / 72 / max(width_in, 0.1) * 1.03))
    return lines


def line_h(size, spacing=1.12):
    # Noto Sans CJK KR 은 줄 높이가 글자 크기의 약 1.4배다
    return size / 72 * 1.38 * spacing


def _set_final(ch, f):
    o = ord(ch) - 0xAC00
    return chr(0xAC00 + (o // 28) * 28 + f)


def polite(t):
    """content.js 의 '~한다' 체 문장을 슬라이드용 '~합니다' 체로 바꾼다."""
    import re as _re

    def repl(m):
        w = m.group(1)
        last = w[-1]
        if w.endswith("는") and len(w) > 1:
            return w[:-1] + "습니다"
        if "\uac00" <= last <= "\ud7a3":
            f = (ord(last) - 0xAC00) % 28
            if f == 4:
                return w[:-1] + _set_final(last, 17) + "니다"
            if f == 0:
                return w[:-1] + _set_final(last, 17) + "니다"
            return w + "습니다"
        return w + "입니다"
    return _re.sub(r"(\S+?)다(?=\.|$| \()", repl, t)


def _font(run, size, bold=False, color=GRAY, italic=False, font=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor.from_string(color)
    f.name = font or LATIN
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
        latin.addnext(ea)
    ea.set("typeface", EA)


def _norm_paras(content):
    """str | [para] 를 [[(text, opts)]] 로 만든다. para 는 str 이거나 run 목록."""
    if isinstance(content, str):
        return [[(line, {})] for line in content.split("\n")]
    out = []
    for p in content:
        if isinstance(p, str):
            out.append([(p, {})])
        else:
            out.append([(r, {}) if isinstance(r, str) else r for r in p])
    return out


def text(slide, x, y, w, h, content, size=20, color=GRAY, bold=False, align="l", anchor="t",
         spacing=1.12, after=4, font=None, check=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    bp = tf._txBody.bodyPr
    for k in ("lIns", "rIns", "tIns", "bIns"):
        bp.set(k, "0")
    bp.set("anchor", {"t": "t", "m": "ctr", "b": "b"}[anchor])
    paras = _norm_paras(content)
    plain = []
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = spacing
        p.space_after = Pt(after)
        p.space_before = Pt(0)
        line = ""
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", False), o.get("font", font))
            line += t
        plain.append(line)
    if check:
        need = sum(est_lines(l, w, size) for l in plain) * line_h(size, spacing) + (len(plain) - 1) * after / 72
        if need > h + 0.05:
            WARN.append(f"slide {len(prs.slides._sldIdLst)}: 글상자 넘침 가능 ({need:.2f}in > {h:.2f}in) '{plain[0][:30]}'")
    return box


def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.8, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def chip(slide, x, y, w, h, label, fill=BLUE, color=WHITE, size=13, bold=True):
    rect(slide, x, y, w, h, fill=fill)
    text(slide, x, y, w, h, label, size=size, color=color, bold=bold, align="c", anchor="m", after=0, check=False)


# ---------- 슬라이드 틀 ----------
def _strip_placeholders(slide, keep=(0,)):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in keep:
            ph._element.getparent().remove(ph._element)


def _set_title(slide, lines, size, color=BLUE):
    tf = slide.shapes.title.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        _font(r, size, False, color)


def _footer(slide):
    text(slide, 2.0, 7.12, 8.8, 0.23, FOOTER, size=10, color=WHITE, after=0, check=False)
    box = slide.shapes.add_textbox(Inches(12.58), Inches(7.12), Inches(0.32), Inches(0.25))
    tf = box.text_frame
    bp = tf._txBody.bodyPr
    for k in ("lIns", "rIns", "tIns", "bIns"):
        bp.set(k, "0")
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{5C7A3F71-3D8E-4C1B-9A44-2E5B7D0F6A11}")
    fld.set("type", "slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"))
    rPr.set("lang", "ko-KR")
    rPr.set("sz", "1100")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", WHITE)
    etree.SubElement(rPr, qn("a:latin")).set("typeface", LATIN)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", EA)
    etree.SubElement(fld, qn("a:t")).text = str(len(prs.slides._sldIdLst))


def content(title, notes, layout="Title Only", size=None):
    s = prs.slides.add_slide(LAYOUT[layout])
    _strip_placeholders(s)
    if size is None:
        size = 31 if est_lines(title, 11.9, 31) == 1 else 27
    _set_title(s, [title], size)
    _footer(s)
    s.notes_slide.notes_text_frame.text = notes
    return s


def cover(layout, title_lines, sub, meta, notes):
    s = prs.slides.add_slide(LAYOUT[layout])
    _strip_placeholders(s)
    _set_title(s, title_lines, 38)
    text(s, 0.72, 4.19, 7.65, 1.75, sub, size=23, spacing=1.12)
    text(s, 0.72, 6.36, 7.4, 0.44, meta, size=12, check=False)
    s.notes_slide.notes_text_frame.text = notes
    return s


def footnote(slide, t):
    text(slide, 0.71, 6.48, 11.9, 0.46, t, size=12.5)


def ref(slide, t):
    text(slide, 10.7, 6.97, 1.9, 0.12, t, size=7.5, align="r", after=0, check=False)


# ---------- 반복 구성 요소 ----------
def columns(slide, items, y=1.43, n=None, head_size=27, body_size=21, colors=None):
    n = n or len(items)
    gap = 0.25
    w = (11.9 - gap * (n - 1)) / n
    colors = colors or [TEAL, LIME, TEAL, LIME]
    for i, (head, body) in enumerate(items):
        x = 0.71 + i * (w + gap)
        rect(slide, x, y, w, 0.1, fill=colors[i % len(colors)])
        text(slide, x, y + 0.42, w, 0.8, head, size=head_size, color=BLUE, bold=True)
        if body:
            text(slide, x, y + 1.4, w - 0.05, 6.3 - (y + 1.4), body, size=body_size)
    return w


def table(slide, x, y, colw, rows, header=True, size=16, hsize=16, rowh=0.5, bold_first=False, mono_cols=()):
    nrows, ncols = len(rows), len(colw)
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * nrows))
    tbl = gf.table
    for j, cw in enumerate(colw):
        tbl.columns[j].width = Inches(cw)
    heights = []
    for i in range(nrows):
        sz = hsize if header and i == 0 else size
        lines = max(est_lines(str(rows[i][j]), colw[j] - 0.24, sz) for j in range(ncols))
        need = lines * line_h(sz, 1.06) + 0.14
        heights.append(max(0.5 if header and i == 0 else rowh, need))
    for i in range(nrows):
        tbl.rows[i].height = Inches(heights[i])
        for j in range(ncols):
            cell = tbl.cell(i, j)
            is_head = header and i == 0
            fill = BLUE if is_head else (LIGHT if (i - (1 if header else 0)) % 2 == 0 else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            cell.margin_left, cell.margin_right = Inches(0.12), Inches(0.1)
            cell.margin_top, cell.margin_bottom = Inches(0.06), Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[i][j]
            tf = cell.text_frame
            tf.word_wrap = True
            for k, line in enumerate(str(val).split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.line_spacing = 1.06
                p.space_after = Pt(0)
                r = p.add_run()
                r.text = line
                _font(r, hsize if is_head else size, is_head or (bold_first and j == 0),
                      WHITE if is_head else GRAY, font=MONO if (j in mono_cols and not is_head) else None)
    total = sum(heights)
    if y + total > 6.45:
        WARN.append(f"slide {len(prs.slides._sldIdLst)}: 표가 각주 영역을 침범할 수 있음 (끝 {y + total:.2f}in)")
    return gf


def numbered(slide, items, x=0.71, y=1.42, w=11.9, size=22, pitch=0.9):
    """템플릿의 남색 번호 사각형 목록."""
    for i, it in enumerate(items):
        yy = y + i * pitch
        rect(slide, x, yy, 0.54, 0.52, fill=BLUE)
        text(slide, x + 0.01, yy + 0.05, 0.51, 0.4, str(i + 1), size=21, color=WHITE, bold=True, align="c", after=0, check=False)
        text(slide, x + 0.81, yy + 0.04, w - 0.81, pitch - 0.06, [it] if isinstance(it, list) else it, size=size)


def steps(slide, items, x=0.71, y=3.09, w=7.3, size=19, pitch=0.78):
    """LAB 화면의 진행 순서 (청록 숫자 + 문장)."""
    for i, it in enumerate(items):
        yy = y + i * pitch
        text(slide, x, yy, 0.33, 0.4, str(i + 1), size=21, color=TEAL, bold=True, check=False)
        text(slide, x + 0.44, yy, w - 0.44, pitch - 0.04, it, size=size)


def bullets(slide, items, x, y, w, size=20, pitch=None, gap=0.18):
    """작은 청록 사각형 글머리표. 항목 높이는 글 길이로 계산한다."""
    yy = y
    for it in items:
        plain = it if isinstance(it, str) else "".join(r if isinstance(r, str) else r[0] for r in it)
        h = pitch or est_lines(plain, w - 0.21, size) * line_h(size) + 0.05
        rect(slide, x, yy + size / 72 * 0.45, 0.07, 0.07, fill=TEAL)
        text(slide, x + 0.21, yy, w - 0.21, h, [it] if not isinstance(it, str) else it, size=size)
        yy += h + gap
    return yy


def flow(slide, items, y=1.8, h=1.4, size=20, sub_size=16):
    n = len(items)
    arrow = 0.17
    w = (11.9 - arrow * (n - 1)) / n
    for i, (head, sub) in enumerate(items):
        x = 0.71 + i * (w + arrow)
        rect(slide, x, y, w, h, fill=LIGHT)
        paras = [[(head, {"bold": True, "color": BLUE, "size": size})]]
        if sub:
            paras.append([(sub, {"size": sub_size, "color": GRAY})])
        text(slide, x + 0.12, y + 0.1, w - 0.24, h - 0.2, paras, size=size, align="c", anchor="m", after=2)
        if i < n - 1:
            text(slide, x + w + 0.005, y + h / 2 - 0.22, arrow, 0.4, "›", size=23, color=TEAL, align="c", after=0, check=False)
    return w


def prompt_box(slide, x, y, w, h, label, body, size=16, highlight=True):
    rect(slide, x, y, w, h, fill=LIGHT, line=LINE)
    text(slide, x + 0.22, y + 0.16, w - 0.44, 0.3, label, size=13, color=TEAL, bold=True, after=0, check=False)
    paras = []
    import re as _re
    for ln in body.split("\n"):
        parts = [p for p in _re.split(r"((?<!\{CP)(?<!\{B)\[[^\]]+\])", ln) if p]
        if not parts:
            parts = [""]
        paras.append([(p, {"bold": True, "color": BLUE}) if (highlight and p.startswith("[")) else (p, {}) for p in parts])
    text(slide, x + 0.22, y + 0.55, w - 0.44, h - 0.68, paras, size=size, spacing=1.08, after=2)


def panel(slide, x, y, w, h, header):
    """템플릿의 화면 도식: 연회색 판 + 남색 머리띠."""
    rect(slide, x, y, w, h, fill=LIGHT, line=LINE)
    rect(slide, x, y, w, 0.5, fill=BLUE)
    text(slide, x + 0.22, y + 0.1, w - 0.44, 0.32, header, size=15, color=WHITE, after=0, check=False)


def lab_overview(lab, goal, steps_list, time_line, ref_t, notes, check_size=16, check_count=None):
    s = content(f"{lab['id']} | {lab['title']}", notes)
    rect(s, 0.71, 1.42, 11.9, 0.75, fill=LIGHT)
    text(s, 0.91, 1.42, 11.45, 0.75, goal, size=21, color=BLUE, bold=True, anchor="m", after=0)
    text(s, 0.71, 2.51, 7.0, 0.45, "진행 순서", size=22, color=BLUE, bold=True, check=False)
    steps(s, steps_list, pitch=0.95)
    checks = [polite(c) for c in lab["checks"][:check_count]]
    size = check_size
    while True:
        hs = [est_lines("□ " + c, 3.68, size) * line_h(size) + 0.02 for c in checks]
        if 3.3 + sum(hs) + 0.1 * len(hs) <= 6.15 or size <= 13:
            break
        size -= 0.5
    box_h = max(3.4, 3.3 + sum(hs) + 0.1 * len(hs) - 2.48 + 0.1)
    rect(s, 8.47, 2.48, 4.14, box_h, fill=LIGHT)
    text(s, 8.7, 2.71, 3.7, 0.44, "완료 확인", size=22, color=BLUE, bold=True, check=False)
    yy = 3.3
    for c, hh in zip(checks, hs):
        text(s, 8.71, yy, 3.68, hh, "□ " + c, size=size)
        yy += hh + 0.1
    if yy > 6.3:
        WARN.append(f"slide {len(prs.slides._sldIdLst)}: 완료 확인 상자 넘침 ({yy:.2f})")
    text(s, 0.71, 6.02, 11.7, 0.32, time_line, size=14, color=TEAL, bold=True, check=False)
    ref(s, ref_t)
    return s


def tree(slide, x, y, w, rows, size=16, pitch=0.42):
    rect(slide, x, y, w, pitch * len(rows) + 0.3, fill=WHITE, line=LINE)
    for i, (lvl, name, detail, bold) in enumerate(rows):
        runs = [(name, {"bold": bold, "color": BLUE})]
        if detail:
            runs.append(("   " + detail, {"color": GRAY, "size": size - 2}))
        text(slide, x + 0.22 + lvl * 0.38, y + 0.15 + i * pitch, w - 0.4 - lvl * 0.38, pitch, [runs], size=size, after=0)


def L(i):
    return K["labs"][i]


def P(lab, n):
    return next(p for p in lab["prompts"] if p["label"] == n)["text"]


# =====================================================================
# 오프닝
# =====================================================================
cover("Title Slide", ["Tosca Commander MCP", "입문 교육"],
      "AI 어시스턴트로 테스트케이스 만들기\nCommander 2026.1 · 스킬팩 1.2.1 기준\n3시간 20분 · 실습 포함",
      "QA 교육 · 실습 자료 v1.0 · 2026년 10월",
      "과정 소개. 오늘은 Tosca Commander에 들어 있는 MCP Server를 써서, AI 어시스턴트에게 말로 요청해 테스트케이스를 만드는 방법을 배운다. 기준 버전은 Commander 2026.1 LTS Patch 1 이상이고, Tricentis가 공개한 commander-mcp 스킬팩 1.2.1 내용을 기준으로 했다. Tosca를 많이 안 써 본 분도 따라올 수 있게 용어부터 짚고 간다.")

s = content("교육을 마치면 할 수 있는 것",
            "오늘 목표는 세 가지다. MCP Server가 Commander와 AI 사이에서 무엇을 하는지 설명할 수 있고, 기존 모듈로 테스트케이스를 만들어 ScratchBook에서 검증할 수 있고, 프롬프트를 네 가지 요소로 써서 템플릿을 업무에 맞게 고칠 수 있으면 된다. 오늘 실습에서 만든 테스트케이스와 프롬프트는 그대로 가져가서 쓴다.",
            layout="Three Content")
columns(s, [
    ("MCP 이해", "Tosca MCP Server가 Commander와 AI 어시스턴트 사이에서 하는 일을 설명합니다."),
    ("테스트케이스 생성", "기존 모듈로 수동·자동 테스트케이스를 만들고 ScratchBook에서 검증합니다."),
    ("템플릿 활용", "네 가지 요소로 프롬프트를 쓰고, 템플릿을 우리 업무에 맞게 고칩니다."),
])
footnote(s, "기준: Tosca Commander 2026.1 LTS Patch 1 이상 · Tricentis commander-mcp 스킬팩 1.2.1")

s = content("3시간 20분 운영 계획",
            "전체 3시간 20분 중 절반 가까이가 실습이다. 이론은 실습에 필요한 만큼만 다룬다. 실습은 Lab 0부터 Lab 5까지 순서대로 진행하는데, 앞 Lab에서 만든 결과를 뒤 Lab에서 쓰기 때문에 건너뛰지 않는다. 휴식은 Lab 3가 끝난 뒤 10분. 그 전까지 95분이므로 Lab 사이사이 진행 상황을 확인한다.")
table(s, 0.71, 1.42, [2.3, 5.2, 4.4], [
    ["경과 시간", "진행 내용", "핵심 결과"],
    ["00:00–00:10", "오프닝 · Tosca 용어", "실습에 필요한 용어 8개"],
    ["00:10–00:35", "Tosca MCP 이해", "할 수 있는 일과 맡기지 않을 일"],
    ["00:35–01:35", "실습 Lab 0~3", "수동·자동 테스트케이스 생성"],
    ["01:35–01:45", "휴식", "10분"],
    ["01:45–02:15", "실습 Lab 4~5", "검증 실행, 실패 분석, 저장"],
    ["02:15–02:45", "활용사례 · 프롬프트 작성", "네 요소로 고쳐 쓰기"],
    ["02:45–03:10", "프롬프트 템플릿", "T00~T18 쓰는 법"],
    ["03:10–03:20", "정리 · Q&A", "다음 단계"],
], size=16.5, hsize=16.5, rowh=0.52)
footnote(s, "설명 90분 / 직접 실습 90분 / 휴식 10분 / 정리 10분 = 200분")

s = content("오늘 쓰는 Tosca 용어 8개",
            "AI가 돌려주는 답에는 Tosca 용어가 그대로 나온다. 용어를 모르면 결과가 맞는지 판단할 수 없다. 8개만 기억하면 오늘 실습은 충분하다. 특히 ActionMode의 Input과 Verify 차이, ScratchBook과 ExecutionList 차이는 실습에서 계속 나오니 짚어 준다.")
table(s, 0.71, 1.42, [2.9, 9.0], [
    ["용어", "뜻"],
    ["Workspace", "작업 공간. 싱글 유저(내 PC)와 멀티 유저(공통 리포지토리 공유)"],
    ["Module", "테스트할 화면의 기술 정보. XScan으로 만들고, 화면 요소 하나는 ModuleAttribute"],
    ["TestCase", "모듈을 순서대로 놓아 만든 테스트 절차"],
    ["TestStep", "테스트케이스의 한 단계. 스텝에 넣는 값은 TestStepValue"],
    ["ActionMode", "값을 쓰는 방식. Input(입력), Verify(검증), Buffer(재사용), WaitOn(대기)"],
    ["ScratchBook", "시험 실행하는 곳. 결과가 정식 기록으로 남지 않음"],
    ["ExecutionList", "정식 실행 단위. 실행 결과(ExecutionLog)가 남음"],
    ["Check out / Check in", "멀티 유저에서 잠그고 고친 뒤 공통 리포지토리에 반영"],
], size=15.5, hsize=16, rowh=0.5, bold_first=True)
footnote(s, "AI 답변에 이 용어가 그대로 나옵니다. 자세한 설명은 워크북 1.3에 있습니다.")
ref(s, "워크북 1.3")

s = content("Tosca에서 테스트가 만들어지는 흐름",
            "Tosca 테스트는 화면을 스캔해 모듈을 만들고, 모듈을 조합해 테스트케이스를 만들고, 실행해서 결과를 보는 순서로 만들어진다. MCP로 AI가 도와줄 수 있는 구간은 2번부터 5번까지다. 2번은 이미 있는 모듈을 찾아 쓰는 것까지다. 1번 화면 스캔은 XScan 화면에서 사람이 해야 하고, AI는 화면을 스캔해 UI 모듈을 만들 수 없다. 그래서 실습용 모듈은 미리 준비해 두었다. 참고로 API 모듈은 MCP 도구로 만들 수 있다.")
bw = flow(s, [("1 화면 스캔", "XScan"), ("2 모듈", "화면 정보 저장"), ("3 테스트케이스", "스텝과 값"),
              ("4 실행", "ScratchBook · ExecutionList"), ("5 결과 분석", "실패 원인")], size=20, sub_size=15)
rect(s, 0.71, 3.4, bw, 0.08, fill=LIME)
text(s, 0.71, 3.58, bw, 0.4, "사람이 직접", size=15, color=GRAY, bold=True, check=False)
rect(s, 0.71 + bw + 0.17, 3.4, 11.9 - bw - 0.17, 0.08, fill=TEAL)
text(s, 0.71 + bw + 0.17, 3.58, 6, 0.4, "AI 어시스턴트에게 말로 요청할 수 있는 구간", size=15, color=TEAL, bold=True, check=False)
text(s, 0.85, 4.3, 11.3, 1.6, "화면 스캔은 사람이 XScan에서 합니다.\n모듈 찾기부터 결과 분석까지는 AI에게 문장으로 요청합니다.", size=24)
footnote(s, "AI는 화면을 스캔해 UI 모듈을 만들지 못합니다. API 모듈은 MCP 도구(create_api_module)로 만들 수 있습니다.")

# =====================================================================
# Part 1
# =====================================================================
cover("Section Header", ["01", "Tosca MCP 이해"],
      "MCP가 무엇인지\nCommander에서 AI가 할 수 있는 일과\n맡기지 않을 일",
      "Tosca Commander MCP 입문 · 실습 자료 v1.0",
      "Part 1은 25분. MCP 자체를 깊게 파지는 않는다. 오늘 실습을 이해하는 데 필요한 정도, 그리고 AI에게 맡기면 안 되는 일이 무엇인지까지가 목표다.")

s = content("MCP는 AI와 프로그램을 잇는 표준 규격",
            "MCP는 Model Context Protocol의 약자다. AI 앱이 외부 프로그램의 기능을 도구처럼 불러 쓸 수 있게 해 주는 공개 규격이다. 예전에는 AI 앱과 프로그램을 잇는 연동을 조합마다 따로 만들어야 했다. MCP가 생긴 뒤로는 각자 규격 하나만 맞추면 된다. Tosca 입장에서는 서버 하나만 만들면 Copilot, Cursor, Claude 어디서든 붙는다는 뜻이다. 2024년 11월 Anthropic이 공개했고, 2025년 12월 Linux Foundation 산하 Agentic AI Foundation으로 넘어갔다.",
            layout="Comparison")
rect(s, 0.71, 1.42, 11.9, 0.75, fill=LIGHT)
text(s, 0.91, 1.42, 11.45, 0.75, "MCP(Model Context Protocol): AI 앱이 외부 프로그램 기능을 '도구'로 쓰게 하는 공개 규격", size=19, color=BLUE, bold=True, anchor="m", after=0)
text(s, 0.71, 2.5, 5.6, 0.55, "연혁과 지원", size=24, color=BLUE, bold=True, check=False)
text(s, 6.77, 2.5, 5.6, 0.55, "Tosca에서의 의미", size=24, color=BLUE, bold=True, check=False)
bullets(s, ["2024년 11월 Anthropic이 공개", "2025년 12월 Linux Foundation 산하 Agentic AI Foundation으로 이관", "Claude, GitHub Copilot, Cursor 등이 지원"], 0.71, 3.2, 5.6, size=20)
bullets(s, ["AI 앱마다 연동을 따로 만들 필요가 없음", "Tosca는 MCP 서버 하나로 여러 AI 앱과 연결", "AI 앱 하나에 여러 MCP 서버를 붙일 수 있음"], 6.77, 3.2, 5.6, size=20)
footnote(s, "Tosca Cloud용 MCP 서버는 별개 제품입니다. 이 과정은 Commander(설치형) 기준입니다.")

s = content("MCP를 이루는 세 가지",
            "Host는 우리가 직접 쓰는 AI 앱이다. VS Code, Cursor, Claude Desktop 같은 것. Client는 Host 안에 들어 있어서 MCP 서버와 통신을 맡는다. 사용자가 따로 신경 쓸 일은 없다. Server는 기능을 도구 형태로 내놓는 쪽이고, 오늘은 Tosca MCP Server가 여기에 해당한다.",
            layout="Three Content")
columns(s, [
    ("Host", "사용자가 직접 쓰는 AI 앱\nVS Code + Copilot, Cursor, Claude Desktop, Claude Code"),
    ("Client", "Host 안에서 MCP 서버와 통신을 맡습니다. 사용자가 따로 다룰 일은 없습니다."),
    ("Server", "기능을 '도구(Tool)'로 내놓는 쪽. 오늘은 Commander 안의 Tosca MCP Server입니다."),
], body_size=20)
footnote(s, "Tosca MCP Server가 제공하는 도구는 스킬팩 1.2.1 기준 32개입니다.")

s = content("요청 하나가 처리되는 순서",
            "다섯 단계 흐름을 실습 내내 보게 된다. 특히 2번, AI가 어떤 도구를 쓸지와 도구에 넣을 값을 스스로 고른다는 점이 중요하다. 그래서 프롬프트가 구체적이어야 한다. 도구가 여러 개 필요한 요청이면 2번에서 4번까지를 여러 번 반복한다.")
flow(s, [("사람", "문장으로 요청"), ("AI", "쓸 도구를 고름"), ("Client", "서버에 도구 호출"),
         ("Server", "Commander에서 실행"), ("AI", "결과를 정리해 답변")], size=21, sub_size=16)
text(s, 0.85, 3.78, 11.3, 1.85, "2번에서 AI가 도구와 입력값을 스스로 고릅니다.\n그래서 경로와 범위를 프롬프트에 구체적으로 적어야 합니다.", size=24)
footnote(s, "도구가 여러 개 필요한 요청이면 2~4번을 반복합니다.")

s = content("Tosca MCP Server는 Commander 안에서 동작합니다",
            "Tosca MCP Server는 별도 프로그램이 아니라 Commander에 들어 있는 애드인이다. 2026.1 LTS Patch 1부터 들어갔고, Patch 2부터는 기본으로 켜져 있다. Commander가 실행되는 동안 내 PC 안에서만 접속되는 주소로 열린다. 그래서 Commander와 워크스페이스가 열려 있어야만 쓸 수 있다. AI는 워크스페이스를 연 사람의 권한으로 움직이고, 멀티 유저라면 체크아웃 체크인 규칙도 똑같이 지켜야 한다. AI 어시스턴트가 쓰는 모델은 대부분 클라우드에서 돌기 때문에, 요청과 결과 일부가 모델 쪽으로 전달된다. 사내 AI 사용 정책을 확인하고 쓰자.",
            size=29)
panel(s, 0.71, 1.42, 6.9, 4.75, "내 PC")
rect(s, 0.95, 2.45, 2.35, 1.7, fill=WHITE, line=LINE)
text(s, 1.1, 2.6, 2.1, 1.4, [[("AI 어시스턴트", {"bold": True, "color": BLUE})], [("VS Code · Cursor · Claude", {"size": 14})]], size=19)
text(s, 3.3, 2.85, 0.9, 0.4, "HTTP", size=14, color=TEAL, bold=True, align="c", check=False)
text(s, 3.3, 3.15, 0.9, 0.5, "⇄", size=24, color=TEAL, align="c", check=False)
rect(s, 4.2, 2.1, 3.2, 3.5, fill=WHITE, line=LINE)
text(s, 4.38, 2.22, 2.9, 0.4, "Tosca Commander 2026.1", size=16, color=BLUE, bold=True, check=False)
rect(s, 4.38, 2.75, 2.84, 0.8, fill=BLUE)
text(s, 4.5, 2.75, 2.6, 0.8, "Tosca MCP Server Addin", size=16, color=WHITE, bold=True, anchor="m", after=0)
text(s, 4.38, 3.6, 2.84, 0.4, "↓", size=18, color=TEAL, align="c", check=False)
rect(s, 4.38, 4.05, 2.84, 1.35, fill=LIGHT)
text(s, 4.52, 4.12, 2.6, 1.25, [[("열린 워크스페이스", {"bold": True, "color": BLUE})], [("싱글: 내 PC 파일", {"size": 13})], [("멀티: 공통 리포지토리", {"size": 13})]], size=16, after=1)
text(s, 0.95, 5.0, 3.2, 0.9, [[("접속 주소", {"bold": True, "color": BLUE, "size": 14})], [(K["meta"]["mcpUrl"], {"font": MONO, "size": 13})]], size=14, check=False)
bullets(s, ["2026.1 LTS Patch 1부터 제공, Patch 2부터 기본으로 켜져 있음",
            "Commander와 워크스페이스가 열려 있을 때만 동작",
            "워크스페이스를 연 사용자 권한으로 동작, 체크인·체크아웃 규칙도 같음",
            "Tosca Cloud MCP와는 별개 제품"], 7.95, 1.5, 4.66, size=18, gap=0.22)
footnote(s, "AI 모델은 대부분 클라우드에서 동작합니다. 요청과 결과 일부가 모델로 전달되므로 사내 AI 사용 정책을 먼저 확인하세요.")

s = content("같은 일을 할 때 달라지는 것",
            "로그인 테스트케이스 하나를 만든다고 해 보자. 직접 하면 폴더 만들고, 모듈을 끌어다 놓고, 값을 하나씩 넣고, ActionMode를 바꾸고, 실행하고, 고치는 일을 전부 클릭으로 한다. MCP를 쓰면 사람이 하는 일은 요청하고, 계획을 확인하고, 결과를 검토하는 세 가지로 줄어든다. 단, 검토를 건너뛰면 안 된다. AI가 고른 모듈이나 값이 틀릴 수 있다.",
            layout="Comparison")
text(s, 0.71, 1.43, 5.6, 0.55, "Commander에서 직접", size=24, color=BLUE, bold=True, check=False)
text(s, 6.77, 1.43, 5.6, 0.55, "MCP로 요청", size=24, color=BLUE, bold=True, check=False)
for i, t in enumerate(["폴더 만들기", "테스트케이스 만들고 이름 바꾸기", "모듈을 6번 끌어다 놓아 스텝 만들기", "값 8개 입력",
                       "검증 스텝의 ActionMode를 Verify로 변경", "ScratchBook에서 실행", "결과 확인 후 틀린 곳 수정"]):
    yy = 2.15 + i * 0.58
    text(s, 0.71, yy, 0.33, 0.4, str(i + 1), size=18, color=TEAL, bold=True, check=False)
    text(s, 1.12, yy, 5.1, 0.55, t, size=18)
for i, (who, t) in enumerate([("사람", "시나리오와 위치를 문장으로 요청"), ("AI", "모듈을 찾아 스텝 계획을 보여줌"), ("사람", "계획 확인 후 승인"),
                              ("AI", "생성, 검증 실행, 결과 보고"), ("사람", "Commander에서 결과 검토")]):
    yy = 2.15 + i * 0.8
    chip(s, 6.77, yy, 0.8, 0.42, who, fill=BLUE if who == "사람" else TEAL, size=14)
    text(s, 7.75, yy + 0.02, 4.8, 0.7, t, size=18, bold=who == "사람")
footnote(s, "클릭은 AI에게 맡기고, 확인은 사람이 직접 합니다.")

s = content("도구 32개, 여섯 묶음",
            "Tosca MCP Server가 내놓는 도구는 스킬팩 1.2.1 기준 32개다. 외울 필요는 없다. AI가 알아서 고른다. 다만 어떤 묶음이 있는지 알면 무엇을 시킬 수 있는지 감이 온다. 실습에서 주로 쓰는 것은 조회, 작성, 실행 묶음이다. 실제 목록은 Commander 패치에 따라 조금 다를 수 있으니, Lab 0에서 AI에게 도구 목록을 직접 물어본다.")
REP = {"read": "get_workspace_info\nlist_assets", "edit": "set_attribute\ndelete_objects", "author": "create_test_cases\nexecute_task",
       "run": "execute_test_suite\nsave_workspace", "api": "create_api_module\nlist_api_connections", "di": "di_connection\nrun_di_sql_statement"}
rows = [["묶음", "개수", "하는 일", "대표 도구"]]
for g in K["toolGroups"]:
    rows.append([g["name"], f"{len(g['tools'])}개", polite(g["summary"]), REP[g["key"]]])
table(s, 0.71, 1.42, [2.05, 0.95, 4.9, 4.0], rows, size=15, hsize=16, rowh=0.62, bold_first=True, mono_cols=(3,))
footnote(s, "commander-mcp 스킬팩 1.2.1 도구 카탈로그 기준. 실제 목록은 Lab 0에서 AI에게 물어 확인합니다.")
ref(s, "워크북 부록 A")

s = content("AI에게 맡기지 않는 일",
            "AI에게 맡기지 않을 일도 미리 정해 둔다. 화면 스캔은 AI가 못 한다. UI 모듈은 사람이 XScan으로 만든다. 비밀번호 같은 민감 정보는 사람이 넣는다. Data Integrity 연결을 AI가 만들면 Allow AI queries 옵션이 꺼진 채로 만들어지고, 켜는 건 사람이 Connection Manager에서 한다. 삭제는 되돌릴 수 없으니 목록 확인 후 따로 승인한다. 실패가 진짜 결함인지 판단하는 것도 사람 몫이다. 테스트를 고쳐서 결함을 덮으면 안 된다.",
            layout="Comparison")
text(s, 0.71, 1.43, 5.6, 0.55, "사람이 직접 한다", size=24, color=BLUE, bold=True, check=False)
text(s, 6.77, 1.43, 5.6, 0.55, "제약과 주의", size=24, color=BLUE, bold=True, check=False)
NF = K["notForAI"]
mk = lambda it: [(it[1] + "  ", {"bold": True, "color": BLUE}), (polite(it[2]), {})]
bullets(s, [mk(NF[0]), mk(NF[1]), mk(NF[3])], 0.71, 2.2, 5.7, size=18, gap=0.25)
bullets(s, [mk(NF[2]), mk(NF[4]), mk(NF[5])], 6.77, 2.2, 5.84, size=18, gap=0.25)
footnote(s, "AI가 만든 Data Integrity 연결은 'Allow AI queries'가 꺼진 채로 생깁니다. 켜는 것은 사람이 Connection Manager에서 합니다.")

s = content("시작 전에 준비할 것",
            "실습 전에 이 네 가지가 준비돼 있어야 한다. Commander는 2026.1 LTS Patch 1 이상이어야 하고, Patch 1이라면 MCP Server 애드인을 직접 켜야 한다. AI 클라이언트는 Tricentis가 테스트했다고 밝힌 네 가지 중 하나를 쓴다. 스킬팩은 Tricentis 문서가 함께 쓰라고 권하고, 스킬팩 안내서에는 MCP 연결과 스킬팩이 둘 다 필요하다고 되어 있다. 이 과정은 설치를 전제로 한다.")
numbered(s, [
    "Tosca Commander 2026.1 LTS Patch 1 이상 (Patch 2부터 MCP Server 애드인 기본 활성화)",
    "실습 워크스페이스 열기 (실습 내내 열어 둡니다. 닫으면 연결이 끊깁니다)",
    "AI 클라이언트 하나: VS Code + GitHub Copilot, Cursor, Claude Code, Claude Desktop",
    "commander-mcp 스킬팩 설치: github.com/Tricentis/mcp-skills 의 Tosca/Commander/MCP",
], size=20, pitch=1.05)
footnote(s, "Tricentis 문서는 스킬팩을 함께 쓰라고 권하고, 스킬팩 안내서는 MCP 연결과 스킬팩이 둘 다 필요하다고 설명합니다.")
ref(s, "워크북 2.1 · 2.3")

s = content("연결 설정은 서버 주소 한 줄",
            "설정은 서버 주소 한 줄이 전부다. 클라이언트마다 등록하는 곳만 다르다. 오른쪽은 VS Code 예시이고, 나머지 클라이언트 설정은 워크북 2장에 그대로 있다. VS Code는 Copilot Chat을 Agent 모드로 바꿔야 도구가 보인다는 점을 꼭 짚는다.",
            layout="Comparison")
table(s, 0.71, 1.42, [2.4, 3.5], [["클라이언트", "등록하는 곳"]] + [[c["name"], c["where"]] for c in K["clients"]],
      size=14.5, hsize=15.5, rowh=0.62, bold_first=True)
text(s, 6.9, 1.42, 5.7, 0.45, "VS Code 예시 · .vscode/mcp.json", size=18, color=BLUE, bold=True, check=False)
rect(s, 6.9, 1.95, 5.71, 2.85, fill=LIGHT, line=LINE)
text(s, 7.15, 2.12, 5.3, 2.6, K["clients"][0]["code"], size=14, font=MONO, spacing=1.0, after=0)
text(s, 6.9, 5.0, 5.7, 1.1, "Copilot Chat을 Agent 모드로 바꾸고, 도구 목록에 tosca-commander가 켜져 있는지 확인합니다.", size=16)
footnote(s, "포트 46248이 기본값입니다. 디버그 빌드는 8080, Commander MCP 설정에서 바꿨다면 그 값을 씁니다.")
ref(s, "워크북 2.2")

# =====================================================================
# Part 2 실습
# =====================================================================
cover("Section Header", ["02", "실습: 테스트케이스 만들기"],
      "Lab 0 연결 확인부터 Lab 5 저장까지\n90분, Lab 3 다음에 휴식 10분",
      "Tosca Commander MCP 입문 · 실습 자료 v1.0",
      "이제 실습. 워크북 3장을 펴게 한다. 각 Lab마다 입력할 프롬프트가 워크북에 그대로 있으니 복사해서 쓰면 된다. [내 이름] 같은 대괄호 부분만 바꾼다. 막히면 Lab 끝의 '막히면' 칸부터 보게 한다. 실습은 처음부터 끝까지 한 대화창에서 이어서 한다.")

s = content("실습 대상: Demo Web Shop 로그인 화면",
            "실습 대상은 Tricentis가 공개한 연습용 쇼핑몰 Demo Web Shop이다. 시나리오는 로그인하고, 계정 이메일이 맞게 보이는지 검증하고, 로그아웃하는 것. 모듈은 미리 XScan으로 스캔해서 워크스페이스에 넣어 두었다. AI는 화면을 스캔해 모듈을 만들 수 없기 때문이다. 실습 계정 이메일은 따로 나눠 주고, 비밀번호는 TCP로 넣어 두었다. 사내 앱으로 바꿔도 흐름은 똑같다.")
rect(s, 0.71, 1.44, 6.55, 4.78, fill=LIGHT, line=LINE)
rect(s, 0.71, 1.44, 6.55, 0.5, fill=BLUE)
text(s, 0.94, 1.54, 6.1, 0.32, "Log in   |   Log out   |   [실습 이메일]", size=15, color=WHITE, after=0, check=False)
text(s, 0.94, 2.2, 6.0, 0.44, "Returning Customer", size=24, color=BLUE, bold=True, check=False)
for i, (lab, val) in enumerate([("Email:", "[실습 이메일]"), ("Password:", "{CP[DWS_Password]}")]):
    yy = 2.95 + i * 0.62
    text(s, 0.95, yy, 1.9, 0.38, lab, size=18, check=False)
    rect(s, 2.85, yy - 0.02, 4.1, 0.42, fill=WHITE, line=LINE)
    text(s, 2.97, yy + 0.02, 3.9, 0.32, val, size=15, check=False)
rect(s, 2.85, 4.35, 2.0, 0.52, fill=BLUE)
text(s, 2.85, 4.35, 2.0, 0.52, "Log in", size=19, color=WHITE, bold=True, align="c", anchor="m", after=0, check=False)
text(s, 0.95, 5.15, 6.0, 0.8, "Login was unsuccessful ...", size=15, color=ORANGE, check=False)
text(s, 7.5, 1.48, 5.1, 0.48, "Login 모듈", size=22, color=BLUE, bold=True, check=False)
for i, a in enumerate(["Email", "Password", "Log in", "Error Message"]):
    text(s, 7.52, 2.05 + i * 0.45, 4.9, 0.4, "→  " + a, size=18, color=BLUE, check=False)
text(s, 7.5, 3.95, 5.1, 0.48, "Header 모듈", size=22, color=BLUE, bold=True, check=False)
for i, a in enumerate(["Log in", "Log out", "Account (로그인 후 이메일)"]):
    text(s, 7.52, 4.52 + i * 0.45, 4.9, 0.4, "→  " + a, size=18, color=BLUE, check=False)
footnote(s, "설명용 도식입니다. Header, Login 모듈은 강사가 XScan으로 미리 스캔했고, OpenUrl·CloseBrowser는 표준 모듈입니다.")
ref(s, "워크북 2.4")

s = content("실습 진행 순서",
            "되돌리기 쉬운 작업부터 한다. Lab 0과 Lab 1은 조회만 하므로 아무것도 바뀌지 않는다. Lab 2와 Lab 3에서 테스트케이스를 만들고, 여기까지 하고 10분 쉰다. Lab 4에서 값을 바꿔 가며 실행하고, Lab 5에서 정리한다. 테스트케이스는 만들 때 바로 저장되므로, 저장을 따로 신경 쓸 곳은 Lab 4의 값 수정이다.")
table(s, 0.71, 1.42, [1.4, 4.3, 1.3, 4.9], [
    ["Lab", "내용", "시간", "워크스페이스 변화"],
    ["Lab 0", "연결 확인", "10분", "없음 (조회)"],
    ["Lab 1", "워크스페이스 둘러보기", "10분", "없음 (조회)"],
    ["Lab 2", "수동 테스트케이스 만들기", "15분", "테스트케이스 생성 (만들 때 바로 저장)"],
    ["Lab 3", "자동 테스트케이스 만들기", "25분", "테스트케이스 생성 (만들 때 바로 저장)"],
    ["휴식", "", "10분", ""],
    ["Lab 4", "ScratchBook 검증 실행과 수정", "20분", "값 수정 (저장 필요), 실행"],
    ["Lab 5", "저장과 정리", "10분", "저장"],
], size=16.5, hsize=16.5, rowh=0.56, bold_first=True)
footnote(s, "실습은 Lab 0부터 Lab 5까지 한 대화창에서 이어서 진행합니다. 새 대화를 열면 AI가 앞의 작업을 기억하지 못합니다.")
ref(s, "워크북 3장")

# ---- Lab 0
lab = L(0)
lab_overview(lab, "AI 어시스턴트가 Commander의 MCP Server에 연결됐는지 확인합니다.",
             ["Tosca Commander를 실행하고 실습 워크스페이스를 엽니다.",
              "AI 클라이언트에서 tosca-commander 서버를 켭니다. VS Code는 Agent 모드로 바꿉니다.",
              "프롬프트 0-1, 0-2를 차례로 입력하고 결과를 기록합니다."],
             "00:35–00:45 / 직접 실습 10분", "워크북 3장 Lab 0",
             "먼저 Commander를 켜고 실습 워크스페이스를 연다. 그다음 AI 클라이언트에서 tosca-commander 서버가 켜져 있는지 본다. 첫 번째 프롬프트로 도구 목록을 받아 보고, 두 번째로 워크스페이스 정보를 받는다. 모드가 single인지 multi인지 기록해 둔다. Lab 5에서 저장 방법이 달라진다.")
s = content("Lab 0 | 입력할 프롬프트와 화면 예시",
            "응답 위에 get_workspace_info 같은 도구 이름이 보이면 MCP로 실제 Commander를 읽은 것이다. 도구 목록이 안 보이면 클라이언트에서 서버가 켜져 있는지, VS Code라면 Agent 모드인지부터 확인한다.",
            layout="Comparison")
prompt_box(s, 0.71, 1.42, 5.75, 1.7, "프롬프트 0-1", P(lab, "0-1"), size=16)
prompt_box(s, 0.71, 3.3, 5.75, 1.7, "프롬프트 0-2", P(lab, "0-2"), size=16)
text(s, 0.71, 5.2, 5.75, 0.8, [[("AI가 쓰는 도구  ", {"bold": True, "color": BLUE}), ("get_workspace_info", {"font": MONO})]], size=16)
panel(s, 6.77, 1.42, 5.84, 4.78, "AI 어시스턴트  ·  tosca-commander 연결됨")
rect(s, 8.0, 2.12, 4.4, 1.05, fill=BLUE)
text(s, 8.15, 2.12, 4.1, 1.05, P(lab, "0-2"), size=13, color=WHITE, anchor="m", after=0)
rect(s, 6.98, 3.35, 5.0, 2.6, fill=WHITE, line=LINE)
text(s, 7.15, 3.47, 4.7, 0.3, "get_workspace_info 실행", size=13, color=TEAL, bold=True, check=False)
text(s, 7.15, 3.85, 4.7, 2.0, "현재 열린 워크스페이스입니다.\n· 이름: MCP_Training_홍길동\n· 경로: C:\\Tosca_Projects\\Tosca_Workspaces\\MCP_Training_홍길동\n· 모드: single (싱글 유저)", size=13.5, after=2)
footnote(s, "화면 예시입니다. 실제 문구와 모양은 클라이언트마다 다릅니다.")
ref(s, "워크북 3장 Lab 0")

# ---- Lab 1
lab = L(1)
lab_overview(lab, "조회만 하는 요청으로 워크스페이스 구조와 실습 모듈을 확인합니다.",
             ["프롬프트 1-1로 TestCases 폴더 구조를 봅니다.",
              "프롬프트 1-2로 실습 모듈과 속성을 확인합니다.",
              "Commander에서 아무것도 바뀌지 않았는지 확인합니다."],
             "00:45–00:55 / 직접 실습 10분", "워크북 3장 Lab 1",
             "이번 Lab은 조회만 한다. 첫 번째 프롬프트로 TestCases 구조를 보고, 두 번째로 실습 모듈과 속성을 확인한다. 결과에 긴 ID가 나와도 신경 쓰지 않는다. 사람은 경로로 이야기하면 된다. Login 모듈에 Email, Password, Log in이 있는지 꼭 확인한다. Lab 3에서 이 속성들을 쓴다.")
s = content("Lab 1 | 입력할 프롬프트와 결과 예시",
            "결과는 오른쪽처럼 트리나 표로 돌아온다. 모양은 달라도 Header, Login 모듈과 속성 이름만 맞으면 된다. 폴더가 없다고 하면 Commander 트리에 보이는 이름 그대로 경로를 다시 적게 한다.",
            layout="Comparison")
prompt_box(s, 0.71, 1.42, 5.75, 1.75, "프롬프트 1-1", P(lab, "1-1"), size=16)
prompt_box(s, 0.71, 3.35, 5.75, 1.75, "프롬프트 1-2", P(lab, "1-2"), size=16)
text(s, 0.71, 5.3, 5.75, 0.8, [[("AI가 쓰는 도구  ", {"bold": True, "color": BLUE}), ("get_object_info, list_assets, get_attributes", {"font": MONO, "size": 14})]], size=16)
panel(s, 6.77, 1.42, 5.84, 4.78, "Tosca Commander  ·  프롬프트 1-2 결과 예시")
tree(s, 6.98, 2.1, 5.42, [
    (0, "Modules", "", True), (1, "MCP실습", "", False), (2, "Header", "", True),
    (3, "Log in · Log out · Account", "", False), (2, "Login", "", True),
    (3, "Email · Password · Log in", "", False), (3, "Error Message", "", False)], size=17, pitch=0.48)
footnote(s, "결과에 나오는 긴 ID(surrogateId)는 몰라도 됩니다. 경로로 이야기하면 됩니다.")
ref(s, "워크북 3장 Lab 1")

# ---- Lab 2
lab = L(2)
lab_overview(lab, "글로 된 테스트 절차를 Tosca 수동 테스트케이스로 옮깁니다.",
             ["[내 이름]을 바꿔 프롬프트 2-1을 입력합니다.",
              "AI가 보여 준 작업 순서를 읽고 2-2로 승인합니다.",
              "Commander에서 테스트케이스를 열어 스텝을 확인합니다."],
             "00:55–01:10 / 직접 실습 15분", "워크북 3장 Lab 2",
             "글로 된 수동 테스트 절차를 그대로 붙여 넣어 Tosca 수동 테스트케이스로 만든다. 프롬프트 마지막 줄 '만들기 전에 순서를 먼저 보여줘'가 중요하다. 작업 상태는 반드시 속성 변경으로 바꾸게 한다. 메뉴 작업으로 바꾸면 대화상자가 떠서 Commander가 멈춘다. 결과는 Commander에서 직접 열어 확인한다.",
             check_size=15)
s = content("Lab 2 | 입력할 프롬프트와 결과",
            "프롬프트 2-1은 글로 된 테스트 절차를 그대로 옮긴 것이다. AI가 작업 순서를 보여 주면 읽어 보고 2-2로 승인한다. 오른쪽처럼 테스트케이스 아래에 수동 스텝 세 개가 생기면 된다.",
            layout="Comparison")
prompt_box(s, 0.71, 1.42, 7.2, 3.9, "프롬프트 2-1", P(lab, "2-1"), size=14)
prompt_box(s, 0.71, 5.45, 7.2, 1.0, "프롬프트 2-2", P(lab, "2-2"), size=13.5)
panel(s, 8.15, 1.42, 4.46, 4.93, "Commander에서 확인")
tree(s, 8.33, 2.1, 4.1, [
    (0, "TestCases / MCP실습 / 홍길동", "", False), (1, "로그인 실패_수동", "", True),
    (2, "1. 사이트 접속", "", False), (2, "2. Log in 링크 클릭", "", False), (2, "3. 틀린 비밀번호로 로그인", "", False)], size=15, pitch=0.5)
text(s, 8.33, 5.0, 4.1, 1.2, "스텝마다 지시와 기대 결과가 들어 있는지 펼쳐서 확인합니다.", size=14)
ref(s, "워크북 3장 Lab 2")

# ---- Lab 3
lab = L(3)
lab_overview(lab, "기존 모듈로 로그인 자동 테스트케이스를 만듭니다. 계획부터 받습니다.",
             ["프롬프트 3-1로 계획만 받습니다. 비밀번호는 {CP[DWS_Password]}로 적습니다.",
              "계획 표의 모듈, 값, ActionMode를 하나씩 확인합니다.",
              "3-2로 승인하고 결과를 Commander에서 확인합니다."],
             "01:10–01:35 / 직접 실습 25분", "워크북 3장 Lab 3",
             "자동화 테스트케이스는 처음부터 만들게 하지 않는다. 먼저 계획만 받는다. 프롬프트에 모듈을 찾을 범위를 정해 주고, 새로 스캔하지 말라고 적는다. 비밀번호는 프롬프트에 쓰지 않고 TCP 참조로 넣는다. 클라우드 모델로 넘어가지 않게 하기 위해서다. 두세 번 고쳐도 안 되면 강사가 준비한 완성본을 복사해 Lab 4로 넘어간다.",
             check_size=15, check_count=4)
s = content("Lab 3 | 프롬프트 3-1: 계획만 받기",
            "프롬프트를 천천히 읽어 준다. 대상(폴더), 작업(자동화 테스트케이스), 조건(모듈 범위, 스캔 금지), 확인(계획을 표로)이 다 들어 있다는 점을 짚는다. Part 3에서 이 네 가지를 다시 다룬다.")
prompt_box(s, 0.71, 1.42, 11.9, 4.85, "프롬프트 3-1", P(lab, "3-1"), size=17)
footnote(s, "비밀번호는 {CP[DWS_Password]} 참조로만 적습니다. 실제 비밀번호를 프롬프트에 쓰지 않습니다.")
ref(s, "워크북 3장 Lab 3")

s = content("Lab 3 | 계획을 확인하고 승인하기",
            "왼쪽처럼 스텝마다 모듈, 속성, 값, ActionMode가 나온 표를 받으면 하나씩 확인한다. 4번 Account 스텝이 Verify인지, Password가 참조로 들어갔는지, 못 찾은 모듈이 없는지가 핵심이다. 맞으면 3-2로 승인한다. 만든 뒤에는 Commander에서 직접 열어 본다. 여기까지 끝나면 10분 휴식.",
            layout="Comparison")
rows = [["#", "모듈", "속성", "값", "ActionMode"]] + K["lab3Steps"]
table(s, 0.71, 1.42, [0.4, 1.55, 1.15, 2.95, 1.35], rows, size=13, hsize=13.5, rowh=0.44, mono_cols=(3,))
prompt_box(s, 8.4, 1.42, 4.21, 1.75, "프롬프트 3-2", P(lab, "3-2"), size=14.5)
text(s, 8.4, 3.42, 4.21, 0.45, "확인할 것", size=20, color=BLUE, bold=True, check=False)
bullets(s, ["4번 Account가 Verify", "Password가 {CP[DWS_Password]} 참조", "맞는 모듈을 못 찾은 스텝 없음"], 8.4, 3.98, 4.21, size=16, gap=0.14)
footnote(s, "모듈이 없다는 답이 오면 그 화면은 사람이 XScan으로 스캔한 뒤 같은 프롬프트로 다시 요청합니다.")
ref(s, "워크북 3장 Lab 3")

cover("Section Header", ["휴식", "10분"],
      "다음: Lab 4 ScratchBook 검증 실행\n방금 만든 '로그인 성공_자동'을 실행합니다.",
      "Tosca Commander MCP 입문 · 실습 자료 v1.0",
      "10분 쉰다. 쉬는 동안 Lab 3이 끝나지 않은 사람은 강사가 완성본 복사를 도와준다. Commander와 AI 대화창은 닫지 않고 그대로 둔다.")

# ---- Lab 4
lab = L(4)
lab_overview(lab, "검증 실행 후, 일부러 낸 실패를 AI와 함께 분석하고 되돌립니다.",
             ["4-1로 저장하고 검증 실행합니다. 모두 통과하면 됩니다.",
              "4-2로 Verify 값을 틀리게 바꾸고 저장 후 다시 실행합니다.",
              "4-3으로 원인만 듣고, 4-4로 되돌려 다시 실행합니다."],
             "01:45–02:05 / 직접 실습 20분", "워크북 3장 Lab 4",
             "만든 테스트케이스를 ScratchBook에서 검증 실행한다. 실행은 저장된 상태로 돌기 때문에, 값을 바꾸면 꼭 저장하고 실행하게 한다. 한 번 통과하면 일부러 Verify 값을 틀리게 바꿔서 실패를 만든다. 그리고 AI에게 원인만 설명하라고 한다. 고치라고 바로 시키지 않는 게 포인트다.",
             check_size=15.5)
s = content("Lab 4 | 통과 → 일부러 실패 → 원인 → 다시 통과",
            "네 개의 프롬프트를 순서대로 입력한다. 브라우저가 열리면 끝날 때까지 손을 뗀다. 4-3에서 '아직 고치지는 마'라고 적는 이유를 설명한다. 분석과 수정을 나눠야 실제 결함을 테스트 수정으로 덮는 일을 막을 수 있다.",
            size=29)
cards = [("4-1", "통과", TEAL, WHITE), ("4-2", "실패", ORANGE, WHITE), ("4-3", "기대값 불일치", LIME, GRAY), ("4-4", "통과", TEAL, WHITE)]
for i, (lbl, res, fill, fc) in enumerate(cards):
    x = 0.71 + (i % 2) * 6.04
    y = 1.42 + (i // 2) * 2.5
    prompt_box(s, x, y, 5.86, 2.32, "프롬프트 " + lbl, P(lab, lbl), size=16)
    chip(s, x + 5.86 - 1.95, y + 0.14, 1.75, 0.36, res, fill=fill, color=fc, size=13)
footnote(s, "실행 중에는 마우스와 키보드를 건드리지 않습니다. 실행은 한 번에 하나, 응답이 끊겨도 다시 실행하지 말고 결과만 다시 묻습니다.")
ref(s, "워크북 3장 Lab 4")

# ---- Lab 5
lab = L(5)
lab_overview(lab, "결과를 저장하고, 멀티 유저에서 더 필요한 일을 확인합니다.",
             ["5-1로 오늘 바꾼 객체를 정리하고 저장합니다.",
              "멀티 유저라면 5-2로 체크인까지 합니다. 오늘은 읽기만 합니다.",
              "Commander를 다시 열어 결과가 남았는지 확인합니다."],
             "02:05–02:15 / 직접 실습 10분", "워크북 3장 Lab 5",
             "마지막으로 오늘 만든 것을 정리하고 저장한다. 테스트케이스는 만들 때 이미 저장됐으니, 여기서 저장하는 건 Lab 4처럼 값을 바꾼 것이다. 싱글 유저는 저장으로 끝이다. 멀티 유저라면 시작할 때 update_all로 최신 내용을 받고, 고칠 객체를 체크아웃하고, 끝에 저장과 체크인을 한다.")
s = content("Lab 5 | 저장과 체크인",
            "체크인은 내가 체크아웃한 객체 전부에 적용된다. 그래서 체크인 전에는 바뀐 객체 목록을 먼저 보는 습관을 들인다. 오른쪽 표는 싱글 유저와 멀티 유저에서 작업 시작과 끝에 할 일이다.",
            layout="Comparison")
prompt_box(s, 0.71, 1.42, 5.75, 1.6, "프롬프트 5-1", P(lab, "5-1"), size=15.5)
prompt_box(s, 0.71, 3.2, 5.75, 2.95, "프롬프트 5-2 (멀티 유저일 때)", P(lab, "5-2").replace("(멀티 유저일 때) ", ""), size=15)
text(s, 6.77, 1.43, 5.84, 0.55, "시작과 끝에 하는 일", size=22, color=BLUE, bold=True, check=False)
table(s, 6.77, 2.1, [1.6, 4.24], [
    ["구분", "할 일"],
    ["싱글 유저", "끝: save_workspace"],
    ["멀티 유저", "시작: update_all → Checkout\n끝: save_workspace → check_in_all"],
], size=15, hsize=15.5, rowh=0.62, bold_first=True)
text(s, 6.77, 4.3, 5.84, 1.8, "check_in_all은 내가 체크아웃한 객체를 전부 체크인합니다. create_test_cases로 만든 테스트케이스는 만들 때 이미 저장·체크인됩니다.", size=16)
ref(s, "워크북 3장 Lab 5")

s = content("실습 중에 막혔을 때",
            "실습에서 자주 막히는 지점이다. 대부분 Commander가 닫혀 있거나, 경로가 틀렸거나, 저장을 안 한 경우다. 그래도 안 되면 AI 응답을 그대로 복사해 보여 달라고 한다. 어떤 도구가 어떤 오류를 냈는지가 가장 빠른 단서다.")
T = K["troubleshooting"]
table(s, 0.71, 1.42, [3.3, 8.6], [["증상", "이렇게 합니다"]] + [[T[i][0], T[i][2]] for i in (0, 2, 3, 4, 5, 7)],
      size=14.5, hsize=15.5, rowh=0.62, bold_first=True)
footnote(s, "그래도 안 되면 AI 응답을 그대로 복사해 강사에게 보여 주세요. 증상 8가지 전체는 워크북 부록 B에 있습니다.")
ref(s, "워크북 부록 B")

# =====================================================================
# Part 3
# =====================================================================
cover("Section Header", ["03", "활용사례와 프롬프트 작성법"],
      "업무에 써 볼 만한 여섯 가지와\n원하는 결과가 나오는 프롬프트 쓰는 법",
      "Tosca Commander MCP 입문 · 실습 자료 v1.0",
      "실습에서 손으로 해 봤으니, 이제 실제 업무에 어디에 쓸지와 프롬프트를 어떻게 써야 원하는 결과가 나오는지 정리한다.")

s = content("업무에 써 볼 만한 여섯 가지",
            "공식 워크플로와 도구를 기준으로 업무에 바로 적용해 볼 수 있는 쓰임새를 여섯 가지로 묶었다. 오늘 실습은 1번과 2번이었다. 4번 실패 원인 분석은 야간 실행 결과를 볼 때 쓸모가 크다. 6번 중 API 모듈 만들기는 대응하는 공식 워크플로가 없어 템플릿이 없으므로 네 요소로 직접 쓴다. Data Integrity는 라이선스와 권한 설정이 필요하다.")
table(s, 0.71, 1.42, [2.45, 5.25, 2.95, 1.25],
      [["활용사례", "이렇게 씁니다", "사람이 확인할 것", "템플릿"]] + [[u["title"], polite(u["body"]), u["human"], u["template"]] for u in K["useCases"]],
      size=13.5, hsize=15, rowh=0.7, bold_first=True)
footnote(s, "API 모듈 만들기는 대응하는 공식 워크플로가 없어 템플릿 없이 네 요소로 직접 씁니다(워크북 4.1 예시).")
ref(s, "워크북 4.1")

s = content("프롬프트 문장이 도구 입력이 됩니다",
            "AI는 프롬프트를 읽고 도구에 넣을 값을 뽑아낸다. 경로를 적으면 create_test_cases의 path가 되고, 모듈 범위를 적으면 list_assets의 검색 범위가 된다. 반대로 프롬프트에 없는 정보는 AI가 추측하거나 다시 묻는다. 추측이 틀리면 엉뚱한 폴더에 테스트케이스가 생긴다.")
rect(s, 0.71, 1.42, 11.9, 1.45, fill=LIGHT)
B_ = lambda t: (t, {"bold": True, "color": BLUE})
M_ = lambda t: (t, {"bold": True, "color": TEAL})
text(s, 0.95, 1.5, 11.4, 1.3, [[B_("TestCases/MCP실습/홍길동"), M_(" ①"), (" 폴더에 ", {}), B_("'로그인 성공_자동'"), M_(" ②"), (" 테스트케이스를 만들어줘. 모듈은 ", {}),
                                 B_("Modules/MCP실습에서만"), M_(" ③"), (" 찾고, Account는 ", {}), B_("Verify"), M_(" ④"), ("로 검증해줘.", {})]],
     size=20, anchor="m", spacing=1.2)
calls = [("③", "list_assets", "assetType = XModule\nroot = Modules/MCP실습"),
         ("AI 판단", "get_object_info", "찾은 모듈의 속성 확인"),
         ("① ②", "create_test_cases", "path = TestCases/MCP실습/홍길동\nname = 로그인 성공_자동"),
         ("④", "set_attribute", "Account 값의\nActionMode = Verify")]
w = (11.9 - 0.17 * 3) / 4
for i, (mark, tool, inp) in enumerate(calls):
    x = 0.71 + i * (w + 0.17)
    rect(s, x, 3.15, w, 2.55, fill=LIGHT)
    text(s, x + 0.16, 3.27, w - 0.32, 0.38, mark, size=16, color=TEAL, bold=True, check=False)
    text(s, x + 0.16, 3.7, w - 0.32, 0.45, tool, size=17, color=BLUE, bold=True, font=MONO, check=False)
    text(s, x + 0.16, 4.25, w - 0.32, 1.35, inp, size=13.5)
    if i < 3:
        text(s, x + w + 0.005, 4.2, 0.17, 0.4, "›", size=23, color=TEAL, align="c", after=0, check=False)
footnote(s, "프롬프트에 없는 정보는 AI가 추측하거나 다시 묻습니다. 추측이 틀리면 엉뚱한 곳에 만들어집니다.")

s = content("프롬프트는 네 가지 요소로 씁니다",
            "Tricentis는 MCP 프롬프트를 쓸 때 구체적으로 쓰고, 맥락을 주고, 단계를 나눠 요청하라고 권한다. 이걸 실무에서 바로 쓰기 좋게 대상, 작업, 조건, 확인 네 가지로 정리했다. 아래 예시처럼 네 가지가 다 들어가면 AI가 추측할 여지가 거의 없다.")
PE = K["promptElements"]
cw = columns(s, [(f"{e['key']}", None) for e in PE], head_size=27)
for i, e in enumerate(PE):
    x = 0.71 + i * (cw + 0.25)
    text(s, x, 2.5, cw, 0.4, e["q"], size=15, color=TEAL, bold=True, check=False)
    text(s, x, 2.95, cw, 1.35, polite(e["desc"]), size=15)
    text(s, x, 4.35, cw, 0.8, e["ex"], size=13.5, color=BLUE)
rect(s, 0.71, 5.3, 11.9, 1.05, fill=LIGHT)
TAG = lambda t: (f"[{t}] ", {"bold": True, "color": TEAL, "size": 13})
text(s, 0.9, 5.3, 11.5, 1.05, [[TAG("대상"), ("TestCases/MCP실습/홍길동 폴더에 ", {}), TAG("작업"), ("'로그인 성공_자동' 테스트케이스를 만들어줘. ", {}),
                                TAG("조건"), ("모듈은 Modules/MCP실습에서만 찾고 새로 스캔하지 마. ", {}), TAG("확인"), ("만들기 전에 스텝 계획을 표로 보여줘.", {})]],
     size=15, anchor="m", after=0)
ref(s, "워크북 4.3")

s = content("같은 요청, 다른 결과",
            "왼쪽은 실제로 자주 보는 요청이다. 짧고 편하지만 AI가 추측해야 할 게 많다. 가운데는 같은 의도를 네 요소로 고쳐 쓴 것이다. 두 번째와 세 번째 예시는 특히 위험한 경우다. 분석과 수정을 한 번에 시키면 결함을 테스트 수정으로 덮을 수 있고, 정리해 달라는 말은 삭제로 이어질 수 있다.")
table(s, 0.71, 1.42, [2.6, 5.9, 3.4], [["아쉬운 요청", "고친 요청", "이유"]] + [[p["bad"], p["good"], polite(p["why"])] for p in K["promptPairs"]],
      size=14, hsize=15.5, rowh=1.3)
ref(s, "워크북 4.5")

s = content("프롬프트 작성 요령 일곱 가지",
            "일곱 가지 중 가장 중요한 것은 세 번째, 바꾸는 일은 계획부터다. 계획을 보면 AI가 무엇을 오해했는지 실행 전에 알 수 있다. 나머지는 계획을 더 정확하게 만들기 위한 요령이다. 공통 머리말 템플릿 T00에 이 요령들을 규칙으로 넣어 두었다.")
numbered(s, [[(t[0] + "   ", {"bold": True, "color": BLUE}), (polite(t[1]), {})] for t in K["tips"]], size=17, pitch=0.68)
footnote(s, "가장 중요한 것은 3번입니다. 일곱 가지 요령은 T00 공통 머리말에 규칙으로 들어 있습니다.")
ref(s, "워크북 4.4")

# =====================================================================
# Part 4
# =====================================================================
cover("Section Header", ["04", "프롬프트 템플릿"],
      "Tricentis 공식 워크플로 18종과\n이 과정에서 만든 프롬프트 템플릿",
      "Tosca Commander MCP 입문 · 실습 자료 v1.0",
      "마지막 파트. 지금 바로 쓸 수 있는 템플릿을 소개한다. 전체 문구는 워크북 5장에 있으니 슬라이드에서는 구성과 쓰는 법만 본다.")

s = content("공식 워크플로 18종과 프롬프트 템플릿",
            "먼저 구분할 것이 있다. Tricentis가 공개한 commander-mcp 스킬팩 1.2.1에 들어 있는 것은 AI가 따라 하는 작업 순서, 즉 워크플로 템플릿 18종이다. 이건 AI용이고 Tricentis가 제공한다. 사람이 입력하는 프롬프트 템플릿 T00~T18은 이 과정에서 그 18종에 하나씩 맞춰 만든 것이다. 번호는 대응 관계를 찾기 쉽게 공식 목록 순서를 따랐다.",
            layout="Three Content")
rect(s, 0.71, 1.42, 11.9, 0.75, fill=LIGHT)
text(s, 0.91, 1.42, 11.45, 0.75, [[("Tricentis 제공  ", {"bold": True, "color": BLUE}), ("AI가 따라 하는 워크플로 18종      ", {}),
                                   ("이 과정에서 작성  ", {"bold": True, "color": BLUE}), ("사람이 입력하는 프롬프트 템플릿 T00~T18", {})]],
     size=17, anchor="m", after=0)
groups = [["조회", "편집", "관리"], ["작성", "확장"], ["실행·분석"]]
colors = [TEAL, LIME, TEAL]
gw = (11.9 - 0.25 * 2) / 3
for ci, gs in enumerate(groups):
    x = 0.71 + ci * (gw + 0.25)
    rect(s, x, 2.45, gw, 0.08, fill=colors[ci])
    yy = 2.68
    for g in gs:
        lst = [t for t in K["templates"] if t["group"] == g]
        text(s, x, yy, gw, 0.4, f"{g}  {len(lst)}", size=18, color=BLUE, bold=True, check=False)
        yy += 0.45
        for t in lst:
            text(s, x, yy, gw, 0.38, [[(t["id"] + "   ", {"bold": True, "color": TEAL}), (t["title"], {})]], size=15, after=0)
            yy += 0.4
        yy += 0.12
footnote(s, "T00 공통 머리말은 대화를 시작할 때 한 번 붙여 넣는 규칙 6줄입니다. 전체 문구는 워크북 5장에 있습니다.")
ref(s, "워크북 5장")

s = content("T00 공통 머리말: 대화를 시작할 때 한 번",
            "T00은 앞에서 본 작성 요령을 규칙으로 바꾼 것이다. 대화를 시작할 때 한 번 붙여 넣으면 이후 요청에 계속 적용된다. 5번 규칙은 Commander가 대화상자 때문에 멈추는 일을 막는다. 스킬팩이 없는 클라이언트에서는 특히 중요하다.")
prompt_box(s, 0.71, 1.42, 11.9, 4.85, "T00 공통 머리말", K["preamble"]["prompt"], size=17, highlight=False)
footnote(s, "스킬팩 지침이 짧게 설치되는 VS Code나 Claude Desktop에서는 T00을 꼭 붙여 넣으세요.")
ref(s, "워크북 5장 T00")

t06 = next(t for t in K["templates"] if t["id"] == "T06")
s = content("템플릿 하나 뜯어보기: T06 자동 테스트케이스 작성",
            "템플릿마다 같은 구성이다. 언제 쓰는지, 프롬프트 본문, 작성 예시, AI가 주로 쓰는 도구, 확인할 점. 본문의 대괄호 부분만 바꿔 쓰면 된다. T06은 Lab 3에서 쓴 프롬프트를 일반화한 것이다. 대응하는 공식 워크플로는 author-automated-test-case다.",
            layout="Comparison", size=28)
prompt_box(s, 0.71, 1.42, 7.1, 4.9, "T06 본문", t06["prompt"], size=15)
yy = 1.42
for head, body, sz, fnt in [("언제 쓰나", t06["when"], 14, None), ("공식 워크플로", t06["workflow"] + ".md", 13, MONO),
                            ("AI가 주로 쓰는 도구", t06["tools"], 12, MONO), ("확인할 점", polite(t06["caution"]), 13, None)]:
    text(s, 8.1, yy, 4.5, 0.36, head, size=16, color=BLUE, bold=True, check=False)
    if fnt == MONO:
        hh = math.ceil(len(body) * 0.62 * sz / 72 / 4.4) * line_h(sz) + 0.05
    else:
        hh = est_lines(body, 4.5, sz) * line_h(sz) + 0.05
    text(s, 8.1, yy + 0.38, 4.5, hh, body, size=sz, font=fnt, check=False)
    yy += 0.38 + hh + 0.22
ref(s, "워크북 5장 T06")

s = content("템플릿 쓰는 순서",
            "실제로 쓸 때는 이 순서를 따른다. 대화를 시작하면 T00 머리말을 먼저 붙여 넣는다. 하려는 일에 맞는 템플릿을 고르고, 대괄호를 채운다. AI가 계획을 보여주면 확인하고 승인한다. 결과를 검토하고 저장한다. 그리고 써 보니 잘 된 프롬프트는 팀 템플릿에 추가해서 같이 쓴다.")
flow(s, [("T00", "붙여 넣기"), ("템플릿", "고르기"), ("[ ]", "채우기"), ("계획", "확인·승인"), ("결과", "검토·저장"), ("팀 템플릿", "추가")], size=20, sub_size=16)
text(s, 0.85, 3.78, 11.3, 1.85, "스킬팩을 설치하면 AI가 요청에 맞는 공식 워크플로를 골라 따릅니다.\n템플릿은 그 워크플로에 필요한 정보를 빠짐없이 넘겨주는 역할입니다.", size=23)
footnote(s, "사내에서 만든 템플릿이 있다면 같은 형식(언제 쓰나, 본문, 예시, 도구, 확인할 점)으로 워크북 5장에 추가합니다.")

# =====================================================================
# 마무리
# =====================================================================
s = content("마무리: 오늘 가져갈 세 가지",
            "오늘 내용을 세 줄로 정리한다. 다음 단계로 사내 앱에 바로 적용해 보길 권한다. 모듈이 준비되지 않은 앱은 MCP로 할 수 있는 게 거의 없으니, 모듈 준비가 첫 번째다. 질문 받는다.")
bullets(s, ["MCP는 AI가 Commander를 다루게 해 주는 연결 규격입니다. Commander 2026.1 Patch 1부터 애드인으로 들어 있습니다.",
            "AI가 만들고, 사람은 계획과 결과를 확인합니다. 화면 스캔, 비밀번호, 삭제, 결함 판단은 사람이 합니다.",
            "템플릿으로 시작해서 우리 업무에 맞게 고칩니다. 잘 된 프롬프트는 팀 템플릿에 추가합니다.",
            "다음 단계: 사내 앱 모듈을 XScan으로 준비하고, 반복 업무 하나에 템플릿을 적용해 봅니다."], 0.71, 1.5, 11.9, size=21, gap=0.3)
footnote(s, "워크북 5장 템플릿과 부록 B 문제 해결은 교육 뒤 업무에서 그대로 쓰면 됩니다.")

s = content("출처와 참고",
            "자료 출처다. 도구 목록과 워크플로 이름은 2026년 10월 기준 commander-mcp 스킬팩 1.2.1을 따랐다. 스킬팩이 갱신되면 템플릿 목록도 함께 고친다.")
table(s, 0.71, 1.42, [4.6, 7.3], [["자료", "주소"]] + [[r[0], r[1]] for r in K["references"]], size=14.5, hsize=15.5, rowh=0.62, mono_cols=(1,))
footnote(s, "도구 목록과 워크플로 이름은 2026년 10월 기준 commander-mcp 스킬팩 1.2.1을 따릅니다.")

# ---------- 저장 ----------
cp = prs.core_properties
cp.title = K["meta"]["course"]
cp.subject = K["meta"]["subtitle"]
cp.author = "QA 교육"
cp.last_modified_by = "QA 교육"
cp.keywords = "Tosca Commander, MCP, AI 어시스턴트, 교육"
import datetime as _dt
cp.created = cp.modified = _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0, tzinfo=None)
cp.revision = 1
prs.save(OUT)

# app.xml 의 슬라이드 수를 실제 값으로 맞춘다
import re as _re
import shutil as _sh
import zipfile as _zip
tmp = OUT + ".tmp"
with _zip.ZipFile(OUT) as zin, _zip.ZipFile(tmp, "w", _zip.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "docProps/app.xml":
            n = str(len(prs.slides)).encode()
            data = _re.sub(rb"<Slides>\d+</Slides>", b"<Slides>" + n + b"</Slides>", data)
            data = _re.sub(rb"<Notes>\d+</Notes>", b"<Notes>" + n + b"</Notes>", data)
        zout.writestr(item, data)
_sh.move(tmp, OUT)
print("written", OUT, len(prs.slides), "slides")
for w in WARN:
    print("WARN", w)
